import argparse
import os
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pyodbc
from openpyxl import load_workbook


UNKNOWN_KEY = -1


@dataclass(frozen=True)
class FactConfig:
    excel_name: str
    fact_table: str
    hours_column: str
    cost_column: str
    revenue_column: str


FACTS = [
    FactConfig("Planned", "FactPlanned", "PlannedHours", "PlannedCost", "PlannedRevenue"),
    FactConfig("Committed", "FactCommitted", "CommittedHours", "CommittedCost", "CommittedRevenue"),
    FactConfig("Actual", "FactActual", "ActualHours", "ActualCost", "ActualRevenue"),
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Load Planned/Committed/Actual Excel files into Construction DWH facts")
    parser.add_argument("--excel-dir", default=".", help="Directory containing Planned.xlsx, Committed.xlsx, Actual.xlsx")
    parser.add_argument(
        "--connection-string",
        default=os.getenv("DWH_CONN_STR"),
        help="ODBC connection string (or set DWH_CONN_STR)",
    )
    parser.add_argument("--schema", default="dwh", help="Target schema name")
    parser.add_argument("--dry-run", action="store_true", help="Validate and resolve keys without inserting data")
    return parser.parse_args()


def scalar(cursor: pyodbc.Cursor, sql: str, params: tuple[Any, ...] = ()) -> Any:
    row = cursor.execute(sql, params).fetchone()
    return row[0] if row else None


def table_exists(cursor: pyodbc.Cursor, schema: str, table: str) -> bool:
    sql = """
    SELECT COUNT(1)
    FROM INFORMATION_SCHEMA.TABLES
    WHERE TABLE_SCHEMA = ?
      AND TABLE_NAME = ?
    """
    return int(scalar(cursor, sql, (schema, table)) or 0) > 0


def column_exists(cursor: pyodbc.Cursor, schema: str, table: str, column: str) -> bool:
    sql = """
    SELECT COUNT(1)
    FROM INFORMATION_SCHEMA.COLUMNS
    WHERE TABLE_SCHEMA = ?
      AND TABLE_NAME = ?
      AND COLUMN_NAME = ?
    """
    return int(scalar(cursor, sql, (schema, table, column)) or 0) > 0


def resolve_dim_key(cursor: pyodbc.Cursor, schema: str, table: str, key_col: str, bk_col: str, bk_value: Any) -> int:
    if bk_value in (None, ""):
        return UNKNOWN_KEY
    sql = f"SELECT TOP 1 {key_col} FROM {schema}.{table} WHERE {bk_col} = ?"
    key = scalar(cursor, sql, (bk_value,))
    return int(key) if key is not None else UNKNOWN_KEY


def resolve_date_key(cursor: pyodbc.Cursor, schema: str, date_key: Any) -> int:
    if date_key in (None, ""):
        return UNKNOWN_KEY
    sql = f"SELECT TOP 1 DateKey FROM {schema}.DimDate WHERE DateKey = ?"
    key = scalar(cursor, sql, (int(date_key),))
    return int(key) if key is not None else UNKNOWN_KEY


def create_audit_row(cursor: pyodbc.Cursor, schema: str, source_system: str) -> int | None:
    if not table_exists(cursor, schema, "DimAudit"):
        return None

    sql = f"""
    INSERT INTO {schema}.DimAudit (BatchId, SourceSystem, LoadStartDtm, LoadStatus)
    OUTPUT inserted.AuditKey
    VALUES (?, ?, SYSDATETIME(), ?)
    """
    batch_id = f"excel_{source_system}"
    row = cursor.execute(sql, (batch_id, source_system, "RUNNING")).fetchone()
    return int(row[0]) if row else None


def complete_audit_row(cursor: pyodbc.Cursor, schema: str, audit_key: int, rows_read: int, rows_inserted: int) -> None:
    sql = f"""
    UPDATE {schema}.DimAudit
    SET LoadEndDtm = SYSDATETIME(),
        LoadStatus = 'SUCCESS',
        RowsRead = ?,
        RowsInserted = ?
    WHERE AuditKey = ?
    """
    cursor.execute(sql, (rows_read, rows_inserted, audit_key))


def fail_audit_row(cursor: pyodbc.Cursor, schema: str, audit_key: int, error_message: str) -> None:
    sql = f"""
    UPDATE {schema}.DimAudit
    SET LoadEndDtm = SYSDATETIME(),
        LoadStatus = 'FAILED',
        ErrorMessage = ?
    WHERE AuditKey = ?
    """
    cursor.execute(sql, (error_message[:2000], audit_key))


def load_fact_file(cursor: pyodbc.Cursor, schema: str, excel_path: Path, config: FactConfig, dry_run: bool) -> tuple[int, int]:
    workbook = load_workbook(excel_path, data_only=True)
    sheet = workbook.active
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [str(h).strip() if h is not None else "" for h in header_row]
    header_index = {name: idx for idx, name in enumerate(headers)}

    required = {
        "ProjectBK",
        "WBSBK",
        "DateKey",
        "ResourceBK",
        "CustomerBK",
        "SupplierBK",
        "ProductBK",
        "DocumentNumber",
        "DocumentLineNumber",
        "Quantity",
        "Hours",
        "Cost",
        "Revenue",
        "SourceSystem",
    }
    missing = sorted(required - set(header_index.keys()))
    if missing:
        raise ValueError(f"{excel_path.name} is missing columns: {', '.join(missing)}")

    include_audit = column_exists(cursor, schema, config.fact_table, "AuditKey")
    audit_key = create_audit_row(cursor, schema, config.excel_name) if include_audit and not dry_run else None

    rows_read = 0
    rows_inserted = 0

    try:
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row is None:
                continue
            values = list(row)
            if all(v in (None, "") for v in values):
                continue

            rows_read += 1
            project_key = resolve_dim_key(cursor, schema, "DimProject", "ProjectKey", "ProjectBK", values[header_index["ProjectBK"]])
            wbs_key = resolve_dim_key(cursor, schema, "DimWBS", "WBSKey", "WBSBK", values[header_index["WBSBK"]])
            date_key = resolve_date_key(cursor, schema, values[header_index["DateKey"]])
            resource_key = resolve_dim_key(cursor, schema, "DimResource", "ResourceKey", "ResourceBK", values[header_index["ResourceBK"]])
            customer_key = resolve_dim_key(cursor, schema, "DimCustomer", "CustomerKey", "CustomerBK", values[header_index["CustomerBK"]])
            supplier_key = resolve_dim_key(cursor, schema, "DimSupplier", "SupplierKey", "SupplierBK", values[header_index["SupplierBK"]])
            product_key = resolve_dim_key(cursor, schema, "DimProduct", "ProductKey", "ProductBK", values[header_index["ProductBK"]])

            document_number = values[header_index["DocumentNumber"]]
            document_line_number = values[header_index["DocumentLineNumber"]] or None
            quantity = values[header_index["Quantity"]] or 0
            hours = values[header_index["Hours"]] or 0
            cost = values[header_index["Cost"]] or 0
            revenue = values[header_index["Revenue"]] or 0

            if dry_run:
                continue

            if include_audit:
                insert_sql = f"""
                INSERT INTO {schema}.{config.fact_table}
                (ProjectKey, WBSKey, DateKey, ResourceKey, CustomerKey, SupplierKey, ProductKey, AuditKey,
                 DocumentNumber, DocumentLineNumber, Quantity, {config.hours_column}, {config.cost_column}, {config.revenue_column})
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """
                cursor.execute(
                    insert_sql,
                    (
                        project_key,
                        wbs_key,
                        date_key,
                        resource_key,
                        customer_key,
                        supplier_key,
                        product_key,
                        audit_key,
                        document_number,
                        document_line_number,
                        quantity,
                        hours,
                        cost,
                        revenue,
                    ),
                )
            else:
                source_system = values[header_index["SourceSystem"]] or config.excel_name
                insert_sql = f"""
                INSERT INTO {schema}.{config.fact_table}
                (ProjectKey, WBSKey, DateKey, ResourceKey, CustomerKey, SupplierKey, ProductKey,
                 Quantity, {config.hours_column}, {config.cost_column}, {config.revenue_column}, SourceSystem)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """
                cursor.execute(
                    insert_sql,
                    (
                        project_key,
                        wbs_key,
                        date_key,
                        resource_key,
                        customer_key,
                        supplier_key,
                        product_key,
                        quantity,
                        hours,
                        cost,
                        revenue,
                        source_system,
                    ),
                )

            rows_inserted += 1

        if include_audit and audit_key is not None and not dry_run:
            complete_audit_row(cursor, schema, audit_key, rows_read, rows_inserted)

    except Exception as exc:
        if include_audit and audit_key is not None and not dry_run:
            fail_audit_row(cursor, schema, audit_key, str(exc))
        raise

    return rows_read, rows_inserted


def main() -> None:
    args = parse_args()
    if not args.connection_string:
        raise ValueError("Connection string is required. Use --connection-string or set DWH_CONN_STR.")

    excel_dir = Path(args.excel_dir).resolve()
    conn = pyodbc.connect(args.connection_string)

    try:
        cursor = conn.cursor()
        totals_read = 0
        totals_inserted = 0

        for config in FACTS:
            fact_exists = table_exists(cursor, args.schema, config.fact_table)
            if not fact_exists:
                print(f"Skipping {config.fact_table}: table not found")
                continue

            excel_path = excel_dir / f"{config.excel_name}.xlsx"
            if not excel_path.exists():
                raise FileNotFoundError(f"Missing Excel file: {excel_path}")

            rows_read, rows_inserted = load_fact_file(cursor, args.schema, excel_path, config, args.dry_run)
            totals_read += rows_read
            totals_inserted += rows_inserted
            print(f"{config.excel_name}: read={rows_read}, inserted={rows_inserted}")

        if args.dry_run:
            conn.rollback()
            print("Dry run completed. No rows inserted.")
        else:
            conn.commit()
            print(f"Load complete. Total rows read={totals_read}, inserted={totals_inserted}")

    finally:
        conn.close()


if __name__ == "__main__":
    main()
